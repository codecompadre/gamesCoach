import io
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

from docx import Document
import pytest
from rich.console import Console

from gamescoach import documents, cli
from gamescoach.chat import ChatCoach
from gamescoach.documents import DocumentError, document_path, import_document, read_document
from gamescoach.models import AthleteProfile
from gamescoach.storage import Store, StoreError


def word_file(path):
    document = Document()
    document.add_heading("Week 1: CrossFit aerobic capacity", 0)
    document.add_paragraph("Monday: 4 x 4 minutes on the rower, 1 minute easy.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Day"
    table.cell(0, 1).text = "Training"
    table.cell(1, 0).text = "Tuesday"
    table.cell(1, 1).text = "5 sets of 3 front squats at RPE 7"
    document.add_paragraph("Wednesday: rest and recovery.")
    document.save(path)
    return path


def test_word_paragraphs_and_tables_preserve_order_and_original_file(tmp_path):
    path = word_file(tmp_path / "Week 1.docx")
    before = path.read_bytes()
    filename, text = read_document(path)
    assert filename == "Week 1.docx"
    assert text.index("Monday") < text.index("Tuesday") < text.index("Wednesday")
    assert "Row 1: Day | Training" in text
    assert "Row 2: Tuesday | 5 sets of 3 front squats at RPE 7" in text
    assert path.read_bytes() == before


def test_nested_and_merged_cells_do_not_duplicate_workout_text(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "20 minutes easy bike"
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Talk test effort"
    path = tmp_path / "merged.docx"
    document.save(path)
    _, text = read_document(path)
    assert text.count("20 minutes easy bike") == 1
    assert "Talk test effort" in text


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_text_and_markdown_support_unicode(tmp_path, encoding, suffix):
    path = tmp_path / ("week" + suffix)
    path.write_text("CrossFit — 5 × 3 squats", encoding=encoding)
    assert read_document(path)[1] == "CrossFit — 5 × 3 squats"


@pytest.mark.parametrize("value,expected", [
    ('"C:\\Users\\Athlete\\Documents\\Week 1.docx"', "/mnt/c/Users/Athlete/Documents/Week 1.docx"),
    ("D:/Training/week.docx", "/mnt/d/Training/week.docx"),
    ("'/tmp/Week 1.docx'", "/tmp/Week 1.docx"),
])
def test_windows_copy_as_path_and_quotes(value, expected):
    assert document_path(value) == Path(expected)


@pytest.mark.parametrize("suffix", [".doc", ".pdf", ".docm"])
def test_unsupported_formats_offer_conversion(tmp_path, suffix):
    with pytest.raises(DocumentError, match=".docx"):
        read_document(tmp_path / ("workouts" + suffix))


def test_empty_missing_and_corrupt_files(tmp_path):
    with pytest.raises(DocumentError, match="path"):
        read_document(tmp_path / "missing.docx")
    with pytest.raises(DocumentError, match="Enter the path"):
        document_path("")
    path = tmp_path / "empty.txt"
    path.write_text(" ")
    with pytest.raises(DocumentError, match="No readable workout text"):
        read_document(path)
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a word file")
    with pytest.raises(DocumentError, match="Cannot read"):
        read_document(path)
    path = tmp_path / "empty.docx"
    Document().save(path)
    with pytest.raises(DocumentError, match="Images and scanned"):
        read_document(path)


def test_file_and_text_limits_do_not_silently_truncate(tmp_path, monkeypatch):
    path = tmp_path / "week.txt"
    path.write_text("CrossFit workouts " * 20)
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 10)
    with pytest.raises(DocumentError, match="10 MB"):
        read_document(path)
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 10000)
    monkeypatch.setattr(documents, "MAX_TEXT_CHARS", 10)
    with pytest.raises(DocumentError, match="Split it"):
        read_document(path)


def test_expanded_word_archive_limit(tmp_path, monkeypatch):
    path = word_file(tmp_path / "week.docx")
    monkeypatch.setattr(documents, "MAX_EXPANDED_BYTES", 10)
    with pytest.raises(DocumentError, match="expanded"):
        read_document(path)


def test_invalid_xml_and_binary_text_are_actionable(tmp_path):
    path = tmp_path / "corrupt.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "<not-closed>")
    with pytest.raises(DocumentError):
        read_document(path)
    path = tmp_path / "binary.txt"
    path.write_bytes(b"workout\x00binary")
    with pytest.raises(DocumentError, match="binary"):
        read_document(path)


def test_import_persists_text_without_private_directory_path(tmp_path):
    path = word_file(tmp_path / "Week 1.docx")
    session = import_document(path, AthleteProfile(goals="CrossFit Quarterfinals"))
    store = Store(tmp_path / "data")
    store.save_session(session)
    assert store.get_session(session.id).to_dict() == session.to_dict()
    assert session.kind == "document" and session.models == {} and not session.demo
    assert str(tmp_path) not in json.dumps(session.to_dict())
    assert "Tuesday" in session.final_plan


def test_both_coaches_receive_exact_imported_word_text(monkeypatch, tmp_path):
    for name in ("OPENAI_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY", "GAMESCOACH_OPENAI_API_KEY", "GAMESCOACH_GEMINI_API_KEY", "GAMESCOACH_COACH_NOTES_FILE"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("OPENAI_API_KEY", "test-openai")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini")
    from gamescoach import chat
    calls = []
    for classname, label in (("OpenAIProvider", "ChatGPT"), ("GeminiProvider", "Gemini")):
        class Coach:
            def __init__(self, key, model, name=label):
                self.name, self.model = name, model
            def generate(self, *, instructions, prompt):
                calls.append((self.name, instructions, prompt))
                return "Example CrossFit coaching reply."
        monkeypatch.setattr(chat, classname, Coach)
    entry = import_document(word_file(tmp_path / "week.docx"), AthleteProfile())
    result = ChatCoach(AthleteProfile(), [entry]).reply("How should I pace Monday's rower intervals?")
    assert [name for name, _, _ in calls] == ["ChatGPT", "Gemini", "ChatGPT"]
    for _, instructions, prompt in calls:
        context = json.loads(prompt.split("\n", 1)[1].split("\nYour turn:", 1)[0])
        assert context["saved_entries"][0]["kind"] == "document"
        assert context["saved_entries"][0]["final_plan"] == entry.final_plan
        assert "untrusted" in instructions
    assert result.context_ids == [entry.id]


@pytest.fixture
def console_env(tmp_path):
    env = os.environ.copy()
    for name in ("OPENAI_API_KEY", "GEMINI_API_KEY", "GOOGLE_API_KEY", "GAMESCOACH_OPENAI_API_KEY", "GAMESCOACH_GEMINI_API_KEY", "GAMESCOACH_COACH_NOTES_FILE"):
        env.pop(name, None)
    env["GAMESCOACH_DATA_DIR"] = str(tmp_path / "data")
    return env


def run_cli(env, *args, input_text=None):
    return subprocess.run([sys.executable, "-m", "gamescoach", *args], env=env, input=input_text, text=True, capture_output=True, timeout=15)


def test_cli_import_history_show_and_chat_auto_reference(console_env, tmp_path):
    path = word_file(tmp_path / "Week 1.docx")
    imported = run_cli(console_env, "import-file", str(path))
    assert imported.returncode == 0, imported.stdout + imported.stderr
    assert "no API calls" in imported.stdout
    store = Store(console_env["GAMESCOACH_DATA_DIR"])
    entry, = store.list_sessions()
    shown = run_cli(console_env, "show", entry.id)
    assert "Row 2: Tuesday" in shown.stdout
    history = run_cli(console_env, "history")
    assert "Local file" in history.stdout
    chat = run_cli(console_env, "chat", "--demo", input_text="CrossFit pacing question\n/quit\n")
    assert chat.returncode == 0, chat.stdout + chat.stderr
    saved = next(item for item in store.list_sessions() if item.kind == "chat")
    assert entry.id in saved.context_ids


def test_file_command_persists_reference_after_restart(console_env, tmp_path):
    path = word_file(tmp_path / "Week 1.docx")
    result = run_cli(console_env, "chat", "--demo", "--context-limit", "0", input_text=f"First CrossFit question\n/file \"{path}\"\n/quit\n")
    assert result.returncode == 0, result.stdout + result.stderr
    store = Store(console_env["GAMESCOACH_DATA_DIR"])
    chat = next(item for item in store.list_sessions() if item.kind == "chat")
    document = next(item for item in store.list_sessions() if item.kind == "document")
    assert chat.context_ids == [document.id] and len(chat.messages) == 2
    resumed = run_cli(console_env, "chat", "--demo", "--resume", chat.id, input_text="/quit\n")
    assert resumed.returncode == 0 and document.id in resumed.stdout
    assert "imported file" in resumed.stdout


def test_menu_import_and_missing_files_keep_navigation_usable(console_env, tmp_path):
    path = tmp_path / "week.txt"
    path.write_text("Monday: 20 minutes easy bike.")
    result = run_cli(console_env, "--demo", input_text=f"4\n\"{path}\"\n1\n/file /missing.docx\n/readfile /missing.txt\n/quit\nq\n")
    assert result.returncode == 0, result.stdout + result.stderr
    assert "Imported week.txt" in result.stdout
    assert "Cannot open the training file" in result.stdout
    assert "Traceback" not in result.stdout + result.stderr
    assert "Returning to the menu" in result.stdout


def test_failed_document_save_rolls_back_active_chat_context(monkeypatch, tmp_path):
    path = word_file(tmp_path / "week.docx")
    store = Store(tmp_path / "data")
    coaches = []
    class RecordingCoach(ChatCoach):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            coaches.append(self)
    monkeypatch.setattr(cli, "ChatCoach", RecordingCoach)
    def fail(entry):
        raise StoreError("Cannot save imported file")
    monkeypatch.setattr(store, "save_session", fail)
    console = Console(file=io.StringIO())
    inputs = iter([f"/file {path}", "/quit"])
    monkeypatch.setattr(console, "input", lambda prompt: next(inputs))
    cli._chat(console, store, demo=True)
    assert coaches[0].context == []
    assert store.list_sessions() == []


def test_bare_file_command_prompts_for_a_path_and_lists_it_in_help(console_env, tmp_path):
    path = word_file(tmp_path / "Week 1.docx")
    result = run_cli(console_env, "chat", "--demo", "--context-limit", "0", input_text=f"/help\n/file\n\"{path}\"\nCrossFit follow-up\n/quit\n")
    assert result.returncode == 0, result.stdout + result.stderr
    assert "/file [PATH]" in result.stdout
    assert "Training file path" in result.stdout
    assert "Unknown chat command" not in result.stdout
    store = Store(console_env["GAMESCOACH_DATA_DIR"])
    entry = next(item for item in store.list_sessions() if item.kind == "document")
    chat = next(item for item in store.list_sessions() if item.kind == "chat")
    assert chat.context_ids == [entry.id]


@pytest.mark.parametrize("cancel", ["", "/quit", "/ quit"])
def test_bare_file_prompt_can_be_cancelled_without_saving(console_env, cancel):
    result = run_cli(console_env, "chat", "--demo", input_text=f"/file\n{cancel}\n/quit\n")
    assert result.returncode == 0
    assert "File import cancelled" in result.stdout
    assert Store(console_env["GAMESCOACH_DATA_DIR"]).list_sessions() == []


def test_file_prompt_handles_eof_and_version_identifies_current_app(console_env):
    result = run_cli(console_env, "chat", "--demo", input_text="/file\n")
    assert result.returncode == 0
    assert "Traceback" not in result.stdout + result.stderr
    version = run_cli(console_env, "--version")
    assert version.returncode == 0
    from gamescoach import __version__
    assert version.stdout.strip() == f"gamescoach {__version__}"
